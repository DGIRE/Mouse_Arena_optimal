import docx, io, sys

src = r"C:\Projects\Repos\Mouse Arena\Analysis\Plume locations\requests\Request draft 2 (agent-ready).docx"
out = r"C:\Projects\Repos\Mouse Arena\Analysis\Plume locations\requests\_request_text.txt"

d = docx.Document(src)
lines = []
for p in d.paragraphs:
    lines.append(p.text)

# also dump tables if any
for ti, t in enumerate(d.tables):
    lines.append(f"\n[TABLE {ti}]")
    for row in t.rows:
        lines.append(" | ".join(c.text for c in row.cells))

with io.open(out, "w", encoding="utf-8") as f:
    f.write("\n".join(lines))

print(f"wrote {out}, {len(d.paragraphs)} paragraphs, {len(d.tables)} tables")
