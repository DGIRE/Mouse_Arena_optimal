import sys
from docx import Document

for path in sys.argv[1:]:
    print("=" * 80)
    print("FILE:", path)
    print("=" * 80)
    doc = Document(path)
    for p in doc.paragraphs:
        t = p.text.strip()
        if t:
            print(t)
    for ti, tbl in enumerate(doc.tables):
        print(f"--- TABLE {ti} ---")
        for row in tbl.rows:
            cells = [c.text.strip() for c in row.cells]
            print(" | ".join(cells))
