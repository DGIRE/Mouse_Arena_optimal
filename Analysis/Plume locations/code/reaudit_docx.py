import sys, os, re
from docx import Document

BASE = r"C:\Projects\Repos\Mouse Arena\Analysis\Plume locations\reports"
files = ["Odor field report.docx", "Signal enhancement report.docx"]
PLACEHOLDERS = ["{{", "}}", "placeholder", "TODO", "XXX", "FIXME", "TBD"]

def alltext(doc):
    parts = [p.text for p in doc.paragraphs]
    for tbl in doc.tables:
        for row in tbl.rows:
            for c in row.cells:
                parts.append(c.text)
    return "\n".join(parts)

for fn in files:
    path = os.path.join(BASE, fn)
    doc = Document(path)
    txt = alltext(doc)
    low = txt.lower()
    print("="*70)
    print(fn, " chars=", len(txt))
    # placeholder tokens
    hits = [tok for tok in PLACEHOLDERS if tok.lower() in low]
    print("  placeholder tokens:", hits if hits else "NONE")
    # keyword probes
    def has(s): return s.lower() in low
    if "Odor field" in fn:
        print("  mentions 'finer bins fail':", has("finer bins fail"))
        print("  mentions 'largest candidate':", has("largest candidate"))
        print("  says L=40 everywhere:", has("l=40 px") or has("resolved to l=40"))
        print("  cites L=15:", "15" in txt, " cites L=10:", "10" in txt)
        print("  mentions Monte-Carlo / MC null:", has("monte-carlo") or has("monte carlo") or has("mc null") or has("mc-null"))
        print("  mentions gini_mc_null p-value / p<:", has("p =") or has("p<") or has("p-value") or has("pvalue"))
        print("  mentions Loc4 ... includes 0 / not significant:", has("loc4") and (has("includes 0") or has("include 0") or has("not significant") or has("crosses 0")))
        print("  mentions NA / degenerate lambda:", has("na") and (has("degenerate") or has("diluted") or has("non-physical") or has("suppress")))
        print("  mentions multiple comparison / no correction:", has("multiple") or has("no correction") or has("per-test") or has("family-wise") or has("no family"))
    else:
        print("  mentions FPR / false-positive:", has("false-positive") or has("false positive") or has("fpr"))
        print("  cites 0.0319 / 0.0284 (fpr):", has("0.032") or has("0.0319"), has("0.028") or has("0.0284"))
        print("  headlines negative/null distal SNR:", has("negative") or has("null") or has("does not improve") or has("degrade") or has("no improvement"))
        print("  cites 87 negative / 27 positive:", has("87") , has("27"))
        print("  mentions top-10 / top 10:", has("top-10") or has("top 10"))
        print("  mentions H5 guard / after <= before:", has("guard") or has("after") and has("before"))
        print("  mentions drift residual validation:", has("residual") or has("drift-window") or has("drift window"))
    # print a few numeric tokens present for spot check
    nums = set(re.findall(r"0\.0\d{3,}", txt))
    print("  sample decimals present:", sorted(list(nums))[:12])
