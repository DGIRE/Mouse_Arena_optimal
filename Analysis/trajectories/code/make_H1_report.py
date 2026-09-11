"""H1 DOCX report. Every number injected from stats.json / trajectory_metrics.json."""
import json, os
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

BASE = r"C:\Projects\Repos\Mouse Arena\Analysis\trajectories"
DATA = os.path.join(BASE, "data")
FIGDIR = os.path.join(BASE, "reports", "figures")
OUT = os.path.join(BASE, "reports", "H1 - encounters vs tortuosity.docx")

stats = json.load(open(os.path.join(DATA, "stats.json")))
metrics = json.load(open(os.path.join(DATA, "trajectory_metrics.json")))
H1 = stats["H1"]
cal = stats["calibration"]
meta = stats["meta"]
diag = stats["encounter_diagnostics"]
prim = H1["primary_n_encounters_vs_tortuosity"]
sec = H1["secondary_frac_above_vs_tortuosity"]
r_nenc_pl = H1["rho_nenc_pathlen"]
r_tort_pl = H1["rho_tort_pathlen"]
r_rate = H1["rho_encounter_rate_vs_tort"]
perloc = H1["per_location_rho_n_encounters_vs_tortuosity"]
aL = H1["anotherLoc_separate"]

def f3(x): return f"{x:.3f}"
def sci(x): return f"{x:.2e}"

doc = Document()

# ---- styling ----
base = doc.styles["Normal"]
base.font.name = "Calibri"
base.font.size = Pt(10.5)

def h(text, lvl):
    p = doc.add_heading(text, level=lvl)
    return p

title = doc.add_heading("Hypothesis 1 - Ethanol encounters vs trajectory tortuosity", level=0)

sub = doc.add_paragraph()
sub.add_run("Odor-guided orientation trajectory analysis | pooled Loc1-6 | unit = trial").italic = True

# ============ GOAL / H1 / H0 ============
h("1. Goal and hypotheses", 1)
doc.add_paragraph(
    "Goal: test whether the number of ethanol encounters a mouse experiences during a trial "
    "relates to how tortuous (winding) its search trajectory is.")
p = doc.add_paragraph()
p.add_run("H1 (prediction): ").bold = True
p.add_run("Across pooled Loc1-6 trials, ethanol encounters are NEGATIVELY associated with "
          "trajectory tortuosity - more encounters imply a straighter (less tortuous) path.")
p = doc.add_paragraph()
p.add_run("H0: ").bold = True
p.add_run("No monotonic association (Spearman rho = 0).")
p = doc.add_paragraph()
p.add_run("Acceptance rule: ").bold = True
p.add_run("Accept H1 only if rho < 0 with a 95% CI excluding 0; otherwise report the honest "
          "direction and significance.")

# ============ METHODS ============
h("2. Methods", 1)
doc.add_paragraph(
    f"Data source: {meta['aggregate_path']}. Trajectory metrics were computed by the shared "
    f"accessor + traj_common ({meta['traj_common_version']}) and saved to "
    f"data/trajectory_metrics.json and data/trajectory_metrics.h5; group-level statistics were "
    f"saved to data/stats.json. All figures and numbers in this report are read from those saved "
    f"result objects (no re-computation of the science).")
doc.add_paragraph(
    f"Interpreter: the project virtual environment python "
    f"(vras/Scripts/python.exe), matplotlib headless (Agg backend). "
    f"Random seed = {meta['seed']} (used for all trial bootstraps, n_boot = {meta['n_boot']}).")
doc.add_paragraph(
    f"Encounter detector: ethdeconv signal on the head clock, k = {cal['k']} (MAD-based), "
    f"threshold = {cal['threshold']:.6g} a.u.; quiet-period false-positive rate = "
    f"{cal['fpr_per_s']:.3e} per second (target max {cal['fpr_target_max']}, "
    f"within target = {cal['fpr_within_target']}).")
doc.add_paragraph(
    "Tortuosity (definition D6): path_length / straight_line on the cleaned body-centroid track, "
    "where path_length = sum of ||delta p|| over kept consecutive samples and straight_line = "
    "||p_end - p_start|| (guard straight_line >= 1 px). Unitless.")
doc.add_paragraph(
    f"Unit of analysis: the trial. Pooled set = Loc1-6, n = {meta['n_pooled']} trials. "
    f"anotherLoc (n = {meta['n_anotherLoc']}, heterogeneous with extreme path lengths/tortuosity) "
    f"is reported separately and excluded from the pooled statistics and the pooled fit. "
    f"Total trials processed = {meta['n_trials_processed']}, skipped = {meta['n_trials_skipped']}.")
doc.add_paragraph(
    "Statistic: Spearman rho with 95% CI by trial-level bootstrap (the tortuosity distribution "
    "is heavy-tailed). Direction is reported before significance throughout.")

# ============ RESULTS ============
h("3. Results", 1)

# --- headline (honest) ---
p = doc.add_paragraph()
p.add_run("Headline: H1 as stated is NOT supported. ").bold = True
p.add_run(
    f"The primary raw association is strongly POSITIVE (Spearman rho = {f3(prim['spearman_rho'])}, "
    f"95% CI [{f3(prim['ci95_boot_trial'][0])}, {f3(prim['ci95_boot_trial'][1])}], "
    f"p = {sci(prim['p'])}, n = {prim['n']}) - i.e. more encounters go with MORE tortuous paths, "
    f"the OPPOSITE of the predicted negative relationship. This raw positive is, however, largely "
    f"a PATH-LENGTH EXPOSURE ARTIFACT; the exposure-normalized encounter measure is null "
    f"(see 3.2). accept_H1 = {prim['accept_H1']}.")

# 3.1 raw + reconcile
h("3.1 Primary (raw) association and reconciliation with the prior", 2)
doc.add_paragraph(
    f"Raw Spearman rho(n_encounters, tortuosity) = {f3(prim['spearman_rho'])} "
    f"[{f3(prim['ci95_boot_trial'][0])}, {f3(prim['ci95_boot_trial'][1])}], p = {sci(prim['p'])} "
    f"(n = {prim['n']}), direction = {prim['direction']}. The prior Plume-locations Task-1 run "
    f"reported rho = +{prim['prior_rho']} (p = {prim['prior_p']}) on the same trials - weak, "
    f"non-significant, but already POSITIVE (opposite the H1 prediction). The present result "
    f"reproduces that positive sign and is far stronger.")
doc.add_paragraph(prim["reconcile_note"], style="Intense Quote")

# 3.2 exposure confound - LEAD with normalized
h("3.2 Exposure confound and the exposure-normalized (null) result", 2)
doc.add_paragraph(H1["exposure_confound_note"])
doc.add_paragraph(
    f"Both variables scale with path length (exposure): "
    f"rho(n_encounters, path_length) = {f3(r_nenc_pl['spearman_rho'])} "
    f"[{f3(r_nenc_pl['ci95_boot_trial'][0])}, {f3(r_nenc_pl['ci95_boot_trial'][1])}], "
    f"p = {sci(r_nenc_pl['p'])}; and "
    f"rho(tortuosity, path_length) = {f3(r_tort_pl['spearman_rho'])} "
    f"[{f3(r_tort_pl['ci95_boot_trial'][0])}, {f3(r_tort_pl['ci95_boot_trial'][1])}], "
    f"p = {sci(r_tort_pl['p'])}. Longer trials accumulate both more encounters and more "
    f"tortuosity, manufacturing a spurious positive correlation between them.")
p = doc.add_paragraph()
p.add_run("Normalized for exposure, the relationship collapses to null: ").bold = True
p.add_run(
    f"rho(encounter_rate = n_encounters/path_length, tortuosity) = {f3(r_rate['spearman_rho'])} "
    f"[{f3(r_rate['ci95_boot_trial'][0])}, {f3(r_rate['ci95_boot_trial'][1])}], "
    f"p = {sci(r_rate['p'])}, n = {r_rate['n']}. The 95% CI spans 0 and the direction is not "
    f"interpretable; there is no evidence of the predicted (or any) monotonic relationship once "
    f"exposure is controlled.")

# 3.3 secondary frac_above (reproduces prior negative)
h("3.3 Secondary measure: fraction of path spent above threshold", 2)
doc.add_paragraph(
    f"An exposure-insensitive alternative - the fraction of kept samples whose aligned ethdeconv "
    f"exceeds threshold - reproduces the prior negative association: "
    f"rho(frac_path_above_threshold, tortuosity) = {f3(sec['spearman_rho'])} "
    f"[{f3(sec['ci95_boot_trial'][0])}, {f3(sec['ci95_boot_trial'][1])}], p = {sci(sec['p'])} "
    f"(n = {sec['n']}), direction = {sec['direction']} (prior rho = {sec['prior_rho']}). "
    f"Straighter (shorter) paths spend a larger fraction of their path in odor. This is the one "
    f"measure consistent with the spirit of the H1 prediction, but note it is a fraction, not the "
    f"raw encounter count that H1 was stated on.")

# 3.4 per-location table
h("3.4 Per-location Spearman rho (n_encounters vs tortuosity)", 2)
tbl = doc.add_table(rows=1, cols=4)
tbl.style = "Light Grid Accent 1"
hdr = tbl.rows[0].cells
hdr[0].text = "Location"; hdr[1].text = "Spearman rho"; hdr[2].text = "p"; hdr[3].text = "n"
for loc, v in perloc.items():
    r = tbl.add_row().cells
    r[0].text = loc
    r[1].text = f3(v["rho"])
    r[2].text = f"{v['p']:.4g}"
    r[3].text = str(v["n"])
doc.add_paragraph(
    "All six pooled locations show the same positive raw sign, consistent with a shared "
    "exposure driver rather than a location-specific effect.")

# 3.5 anotherLoc separate
h("3.5 anotherLoc (reported separately, excluded from pooled stats/figures)", 2)
aE = aL["n_encounters_vs_tortuosity"]; aF = aL["frac_above_vs_tortuosity"]
doc.add_paragraph(
    f"anotherLoc (n = {aE['n']}): rho(n_encounters, tortuosity) = {f3(aE['rho'])} "
    f"(p = {aE['p']:.4g}); rho(frac_above, tortuosity) = {f3(aF['rho'])} (p = {aF['p']:.4g}). "
    f"These are small-n and not pooled with Loc1-6.")

# 3.6 encounter diagnostics - noise floor (honest)
h("3.6 Encounter diagnostics: the calibrated threshold sits at the noise floor", 2)
amp = diag["n_encounters_amp"]
doc.add_paragraph(
    f"The calibrated encounter threshold ({sci(diag['calibrated_threshold'])} a.u. on "
    f"{diag['signal']}) sits at the deconvolved NOISE FLOOR, not at a physically meaningful odor "
    f"amplitude. The pooled per-trial median ethdeconv (the noise floor) is "
    f"{sci(diag['pertrial_ethd_median_median'])} a.u., and the median ethdeconv AT the "
    f"{diag['n_onsets_pooled']} detected onsets is {sci(diag['onset_ethd_median'])} a.u. - the "
    f"same order of magnitude as both the threshold and the noise floor - whereas real odor reaches "
    f"~{diag['signal_max_ref']:g} a.u. (signal_max_ref). Only a fraction "
    f"{sci(diag['frac_onsets_above_0p01'])} of onsets (i.e. ~0%) exceed the physically meaningful "
    f"{diag['secondary_amp_threshold']:g} a.u. amplitude.")
p = doc.add_paragraph()
p.add_run("At a physically meaningful amplitude, encounters collapse to ~0-1 per trial. ").bold = True
p.add_run(
    f"Re-counting encounters with a secondary amplitude threshold of "
    f"{diag['secondary_amp_threshold']:g} a.u. gives per-trial n_encounters_amp median "
    f"{amp['median']:.0f} (min {amp['min']:.0f}, max {amp['max']:.0f}, mean {amp['mean']:.2f}) - "
    f"i.e. discrete high-amplitude odor contacts are essentially absent, matching the request's "
    f"original guess of roughly 0-1 real contacts per trial.")
doc.add_paragraph(
    f"Interpretation: the ~{cal['n_encounters_distribution_pooled']['median']:.0f} onsets/trial "
    f"index low-amplitude noise/exposure crossings, NOT discrete high-amplitude odor contacts. This "
    f"REINFORCES the H1 conclusion: the raw positive rho(n_encounters, tortuosity) = "
    f"{f3(prim['spearman_rho'])} is a path-length exposure artifact, and the exposure-normalized "
    f"association is null (rho = {f3(r_rate['spearman_rho'])}, p = {sci(r_rate['p'])}, CI spans 0).")
doc.add_paragraph(cal["k_selection_note"], style="Intense Quote")

# --- Figure F1 ---
h("Figure F1 - Encounters vs tortuosity scatter", 2)
f1png = os.path.join(FIGDIR, "F1_encounters_vs_tortuosity.png")
doc.add_picture(f1png, width=Inches(6.2))
doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
cap = doc.add_paragraph()
cap.add_run("Figure F1. ").bold = True
cap.add_run(
    f"One point per pooled Loc1-6 trial (n = {prim['n']}); x = ethanol encounters (counts), "
    f"y = tortuosity (unitless, log scale). Points colored by end location; dashed red line = "
    f"OLS trend of tortuosity on n_encounters (display only, D12). anotherLoc trials (triangles, "
    f"n = {meta['n_anotherLoc']}) are excluded from the fit. Annotation: Spearman "
    f"rho = {f3(prim['spearman_rho'])}, 95% CI [{f3(prim['ci95_boot_trial'][0])}, "
    f"{f3(prim['ci95_boot_trial'][1])}], p = {sci(prim['p'])}. The positive slope is the raw "
    f"(exposure-confounded) association; see 3.2 for the null normalized result.")

# --- Figure F2 ---
h("Figure F2 - Most/least tortuous example paths", 2)
f2png = os.path.join(FIGDIR, "F2_example_paths.png")
doc.add_picture(f2png, width=Inches(6.3))
doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
cap = doc.add_paragraph()
cap.add_run("Figure F2. ").bold = True
cap.add_run(
    "Six example pooled trials: the three most tortuous (left column) and three least tortuous "
    "(right column). Each cell: top = cleaned body-centroid track within the arena box "
    "([0,580] x [0,280] px, equal aspect) with the source endpoint marked (red/white x) and "
    "ethanol contacts (head positions at detected onsets) scatter-colored by ethdeconv at onset "
    "(jet colormap); bottom = that trial's ethdeconv time series on the head clock with the "
    f"detector threshold ({cal['threshold']:.2e} a.u., dashed red) and vertical onset marks. "
    "Each panel is titled with its file_name and tortuosity.")

# ============ CAVEATS ============
h("4. Caveats", 1)
cav = doc.add_paragraph(style="List Bullet")
cav.add_run(
    f"Exposure confound (primary caveat): raw n_encounters and tortuosity both scale with path "
    f"length (rho = {f3(r_nenc_pl['spearman_rho'])} and {f3(r_tort_pl['spearman_rho'])} "
    f"respectively), so the raw positive association is largely an artifact. The "
    f"exposure-normalized encounter_rate vs tortuosity is null "
    f"(rho = {f3(r_rate['spearman_rho'])}, p = {sci(r_rate['p'])}, CI spans 0).")
doc.add_paragraph(cal["low_dynamic_range_note"], style="List Bullet")
doc.add_paragraph(
    f"Encounter-count magnitude: pooled per-trial n_encounters median = "
    f"{cal['n_encounters_distribution_pooled']['median']:.0f} "
    f"(min {cal['n_encounters_distribution_pooled']['min']:.0f}, "
    f"max {cal['n_encounters_distribution_pooled']['max']:.0f}), limiting dynamic range.",
    style="List Bullet")
doc.add_paragraph(
    "Single cohort; infrared/uncalibrated ethdeconv signal in arbitrary units (a.u.); "
    "spatial units in pixels (uncalibrated to physical distance).", style="List Bullet")
doc.add_paragraph(
    "Conclusion: H1 (encounters -> straighter) is NOT supported. The strong raw positive is an "
    "exposure artifact; the exposure-normalized relationship is null; only the frac-above-threshold "
    "measure shows a negative association consistent with the prediction's spirit.", style="List Bullet")

# ============ REPRODUCIBILITY FOOTER ============
h("5. Reproducibility", 1)
doc.add_paragraph(
    f"Seed = {meta['seed']}; traj_common {meta['traj_common_version']}; "
    f"created_utc (result objects) = {meta['created_utc']}.")
doc.add_paragraph("Result files (read-only inputs to this report):")
for fn in ("data/stats.json", "data/trajectory_metrics.json", "data/trajectory_metrics.h5"):
    doc.add_paragraph(fn, style="List Bullet")
doc.add_paragraph("Exact commands (interpreter = vras/Scripts/python.exe):")
mono = doc.add_paragraph()
run = mono.add_run(
    'AR_PY="C:/Projects/Repos/Agentic Research/Research Setup/vras/Scripts/python.exe"\n'
    '"$AR_PY" code/make_F1.py\n'
    '"$AR_PY" code/make_F2.py\n'
    '"$AR_PY" code/make_H1_report.py')
run.font.name = "Consolas"; run.font.size = Pt(9)

doc.save(OUT)
print("Saved:", OUT, "size:", os.path.getsize(OUT))
